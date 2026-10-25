import pymupdf
import numpy as np
import pandas as pd
from pathlib import Path
from openpyxl.styles import Font, PatternFill

header_font = Font(name="Calibri", size=13, bold=True, color="FFFFFF")
header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid") # Classic Navy
category_font = Font(name="Calibri", size=11, bold=True, color="000000")
data_font = Font(name="Calibri", size=11, color="000000")

def is_number(string: str) -> bool:
    string = string.strip()

    string = string.replace(" ", "").replace(",", "")
    try:
        float(string)
        return True
    except ValueError:
        return False

def convert_to_number(cell: str):
    cell = str(cell).strip()

    if cell in ("", "-", "N/A", "n/a"):
        return np.nan
    try:
        value = float(cell.strip("()").replace(",", ""))
    except ValueError:
        return np.nan
    return value


folder_path = Path.cwd().joinpath("TT_Economy_Data/ROE-2026-for-web-1.pdf")
doc = pymupdf.open(folder_path)
y_tol = 15

page = doc[206]
#pix = page.get_pixmap(matrix=fitz.Matrix(2.0, 2.0))

tables_data = page.get_text("blocks")
tables_data = sorted(tables_data, key= lambda b:(b[5], b[1], b[0]))

for index, block in enumerate(tables_data):
    if "Oct" in block[4].strip():
        header_start_index = index
        break

header_row = []
current_col_text = ""
for index, block in enumerate(tables_data[header_start_index:]):
    if not ("Oct" in block[4].strip()) and not ("Sep" in block[4].strip()):
        break_index = header_start_index + index
        break
    else:
        header_row.append(block[4].strip())
    # current_col_text = ""

header_row.insert(0, "")


# GET BODY BLOCKS
body_blocks = []
for block in tables_data[break_index:]:
    if "Source" in block[4].strip():
        break
    else:
        body_blocks.append(block[4].strip())

body_blocks = [item for item in body_blocks if item != ""]
body_blocks[0] = body_blocks[0].replace("\n", "").replace("1", "")
body_blocks[2] = body_blocks[2].replace("\n", "").replace("1", "")
for index, item in enumerate(body_blocks):
    if "Natural Gas Production" in item or "Natural Gas Utilisation" in item:
        label_term = body_blocks.pop(index)
        body_blocks[index] = label_term + " \n" + body_blocks[index]

body_rows =[]
for block in body_blocks:
    block_list = block.split("\n")
    body_rows.append(block_list)

cleaned_body_rows = []
for row in body_rows:
    cleaned_row = [convert_to_number(item.strip()) if is_number(item.strip()) or item.strip() == "n/a" else item.strip() for item in row]
    cleaned_body_rows.append(cleaned_row)

cleaned_rows_df = pd.DataFrame(cleaned_body_rows, columns=header_row)

with pd.ExcelWriter("TT_Nat_Gas_Prod_Util_2026.xlsx", engine="openpyxl") as writer:
    sheet_tab_name = "TT_Nat_Gas_Prod"
    cleaned_rows_df.to_excel(writer, sheet_name=sheet_tab_name, index= False)
    num_format = "#,###.0"
    active_sheet = writer.sheets[sheet_tab_name]

    for row in active_sheet.iter_rows(min_row=2, min_col=2):
        for cell in row:
            cell.number_format = num_format
            cell.font = data_font

    for cell in active_sheet[1]:
        cell.font = header_font
        cell.fill = header_fill

    bold_rows = [2, 3]
    for row_num in bold_rows:
        for cell in active_sheet[row_num]:
            cell.font = Font(bold=True)

    for col in active_sheet.columns:
        max_len = 0
        col_letter = col[0].column_letter

        for cell in col:
            if cell.value is not None:
                val_str = str(cell.value)
                if len(val_str) > max_len:
                    max_len = len(val_str)

        active_sheet.column_dimensions[col_letter].width = max(max_len + 4, 12)