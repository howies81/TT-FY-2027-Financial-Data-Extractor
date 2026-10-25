import fitz

import numpy as np
import pandas as pd
from pathlib import Path
from openpyxl.styles import Font, PatternFill

header_font = Font(name="Calibri", size=13, bold=True, color="FFFFFF")
header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid") # Classic Navy
category_font = Font(name="Calibri", size=11, bold=True, color="000000")
data_font = Font(name="Calibri", size=11, color="000000")

def is_number(string: str) -> bool:
    string = string.strip()

    string = string.replace(" ", "").replace(",", "")
    try:
        float(string)
        return True
    except ValueError:
        return False

def convert_to_number(cell: str):
    cell = str(cell).strip()

    if cell in ("", "-", "N/A"):
        return np.nan
    try:
        value = float(cell.strip("()").replace(",", ""))
    except ValueError:
        return np.nan
    return value


#reader = easyocr.Reader(['en'], gpu=False)
folder_path = Path.cwd().joinpath("TT_Economy_Data/ROE-2026-for-web-1.pdf")
doc = fitz.open(folder_path)
y_tol = 15

page = doc[223]
#pix = page.get_pixmap(matrix=fitz.Matrix(2.0, 2.0))
tables_data = page.get_text("blocks")
tables_data = sorted(tables_data, key= lambda b:(b[5], b[1], b[0]))
# print(tables_data)
# breakpoint()

#Convert pixel map into numpy array

# img = np.frombuffer(pix.samples, dtype= np.uint8).reshape(pix.height, pix.width, pix.n)
# gray = cv2.cvtColor(img, cv2.COLOR_RGB2GRAY)
# processed_img = cv2.adaptiveThreshold(gray, 255, cv2.ADAPTIVE_THRESH_GAUSSIAN_C, cv2.THRESH_BINARY, 11, 2)

for index, block in enumerate(tables_data):
    if "Oct" in block[4].strip():
        header_start_index = index
        break

header_row = []
current_col_text = ""
for index, block in enumerate(tables_data[header_start_index:]):
    if not ("Oct" in block[4].strip()) and not ("Sep" in block[4].strip()):
        break_index = header_start_index + index
        break
    else:
        header_row.append(block[4].strip())
    # current_col_text = ""
cleaned_header_row = []
for column in header_row:
    if "\n" in column:
        split_column = column.split("\n")
        for section in split_column:
            cleaned_header_row.append(section.strip())
    else:
        cleaned_header_row.append(column)


cleaned_header_2nd_row = []
for item in cleaned_header_row:
    if "Sep" in item:
        cleaned_header_2nd_row.append(item)
        cleaned_header_row.remove(item)
    else:
        continue


full_cleaned_header_row = []
item_stitch = ""

for i in range(len(cleaned_header_row)):
    item_stitch += cleaned_header_row[i] + "\n" + cleaned_header_2nd_row[i]
    full_cleaned_header_row.append(item_stitch)
    item_stitch = ""

full_cleaned_header_row.insert(0, "")


##################################################################################
############## LABEL AND DATA ROW COMPILATION #########################################
##################################################################################

# GET BODY BLOCKS
body_blocks = []
for block in tables_data[break_index:]:
    if "Source" in block[4].strip():
        break
    else:
        body_blocks.append(block[4].strip())
body_rows =[]
for block in body_blocks:
    block_list = block.split("\n  \n")
    body_rows.append(block_list)

body_table_1 = body_rows[0]
body_table_1_cleaned = []
for item in body_table_1:
    cleaned_item = item.split("\n")
    cleaned_item = [item.strip() for item in cleaned_item if item != " "]
    if len(cleaned_item) > 7:
        extra_line = cleaned_item.pop(1)
        cleaned_item[0] += " " + extra_line
    #cleaned_item = [np.nan if item == "-" else item for item in cleaned_item]
    cleaned_item = [convert_to_number(item) if is_number(item) or item == "-" else item for item in cleaned_item ]
    body_table_1_cleaned.append(cleaned_item)




body_table_2 = body_rows[1]
# print(body_table_2)
# breakpoint()
body_table_2_cleaned = []
for item in body_table_2:
    cleaned_item = item.split("\n")
    cleaned_item = [item.strip() for item in cleaned_item if item != " "]
    for index, sub_item in enumerate(cleaned_item):
        if "of which:" in sub_item:
            cleaned_item[(index + 1):(index + 1)] = ["-"] * 6
            break
    last_cut_idx = 0
    for index, sub_item in enumerate(cleaned_item):
        if (index + 1) >= len(cleaned_item):
            continue
        elif (is_number(sub_item) or sub_item == "-") and (not is_number(cleaned_item[index + 1]) and cleaned_item[index + 1] != "-") and ((index + 1) < len(cleaned_item)):
            split_row_1 = cleaned_item[last_cut_idx:(index + 1)]
            #split_row_1 = [np.nan if item == "-" else item for item in split_row_1]
            split_row_1 = [convert_to_number(item) if is_number(item) or item == "-" else item for item in split_row_1 ]
            body_table_2_cleaned.append(split_row_1)
            last_cut_idx = (index + 1)

    if last_cut_idx < len(cleaned_item):
        final_row = cleaned_item[last_cut_idx:]
        #final_row = [np.nan if item == "-" else item for item in final_row]
        final_row = [convert_to_number(item) if is_number(item) or item == "-" else item for item in final_row ]
        body_table_2_cleaned.append(final_row)



body_table_3 = body_rows[2]
body_table_3_cleaned = []
for item in body_table_3:
    cleaned_item = item.split("\n")
    cleaned_item = [item.strip() for item in cleaned_item if item != " "]
    cleaned_item = [convert_to_number(item) if is_number(item) or item == "-" else item for item in cleaned_item ]
    body_table_3_cleaned.append(cleaned_item)



body_table_1_df = pd.DataFrame(body_table_1_cleaned, columns=full_cleaned_header_row)
body_table_2_df = pd.DataFrame(body_table_2_cleaned, columns=full_cleaned_header_row)
body_table_3_df = pd.DataFrame(body_table_3_cleaned, columns=full_cleaned_header_row)
full_body_table_df = pd.concat([body_table_1_df, body_table_2_df, body_table_3_df], axis=0)
print(full_body_table_df.to_string(index=False))
# print(data_rows_df.to_string(index= False))
# print(label_data_rows_df.to_string(index= False))


#print(label_data_rows_df.to_string(index= False))
with pd.ExcelWriter("TT_Central_Gov_2026_Financing.xlsx", engine="openpyxl") as writer:
    sheet_tab_name = "Cen_GovTT_Financing"
    full_body_table_df.to_excel(writer, sheet_name=sheet_tab_name, index= False)
    num_format = "#,###.0"
    active_sheet = writer.sheets[sheet_tab_name]

    for row in active_sheet.iter_rows(min_row=2, min_col=2):
        for cell in row:
            cell.number_format = num_format
            cell.font = data_font

    for cell in active_sheet[1]:
        cell.font = header_font
        cell.fill = header_fill

    bold_rows = [2, 3, 7]
    for row_num in bold_rows:
        for cell in active_sheet[row_num]:
            cell.font = Font(bold=True)

    for col in active_sheet.columns:
        max_len = 0
        col_letter = col[0].column_letter

        for cell in col:
            if cell.value is not None:
                val_str = str(cell.value)
                if len(val_str) > max_len:
                    max_len = len(val_str)

        active_sheet.column_dimensions[col_letter].width = max(max_len + 4, 12)

    






